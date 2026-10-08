# -*- coding: utf-8 -*-
"""
build_docx.py —— 把 paper/论文正文.md 导出为 Word 文档（含封面、目录域、正文、参考文献）

用法：  python code/build_docx.py
输出：  paper/论文正文.docx

支持的 Markdown 子集：`#`~`####` 标题、`**加粗**`、`| 表格 |`、`- 列表`、`> 引用`、普通段落。
目录使用 Word 的 TOC 域，打开文档后按 F9 或"右键→更新域"即可生成页码。
"""
from __future__ import annotations

import os
import re

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from utils import ROOT


def set_ea(run, name: str) -> None:
    """安全设置中文字体（兼容 rPr 尚未创建的情况）。"""
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), name)


MD = os.path.join(ROOT, "paper", "论文正文.md")
OUT = os.path.join(ROOT, "paper", "论文正文.docx")

TITLE = "企业数字化转型的文本测度构建及其对全要素生产率的影响"
SUBTITLE = "——基于 2013—2023 年 A 股上市公司年报文本的 NLP 分析"
AUTHOR_LINE = "模拟毕业论文（预演论文）"


def set_cjk_font(doc: Document) -> None:
    """把 Normal 样式设为中文宋体 / 西文 Times New Roman，小四号。"""
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    set_ea(style, "宋体")
    pf = style.paragraph_format
    pf.line_spacing = 1.5
    pf.first_line_indent = Pt(24)


def add_cover(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITLE)
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    if r._element.rPr is None:
        r._element.get_or_add_rPr()
    set_ea(r, "黑体")

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(SUBTITLE)
    r2.font.size = Pt(14)
    if r2._element.rPr is None:
        r2._element.get_or_add_rPr()
    set_ea(r2, "楷体")

    for _ in range(8):
        doc.add_paragraph()

    for line in [AUTHOR_LINE, "学科方向：大数据管理与应用", "完成日期：2026 年 10 月"]:
        pp = doc.add_paragraph()
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = pp.add_run(line)
        rr.font.size = Pt(13)
        set_ea(rr, "宋体")
    doc.add_page_break()


def add_toc(doc: Document) -> None:
    h = doc.add_paragraph()
    r = h.add_run("目  录")
    r.font.size = Pt(16)
    r.font.bold = True
    set_ea(r, "黑体")
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
    run = OxmlElement("w:r")
    txt = OxmlElement("w:t")
    txt.text = "（请在 Word 中按 F9 或右键“更新域”生成目录）"
    run.append(txt)
    fld.append(run)
    p._p.append(fld)
    doc.add_page_break()


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_sep(line: str) -> bool:
    return bool(re.fullmatch(r"\|[\s:\-\|]+\|", line.strip()))


def add_formatted_runs(par, text: str) -> None:
    """处理 **加粗** 与 `代码`。"""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"
        else:
            par.add_run(part)


def build() -> None:
    with open(MD, encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()
    set_cjk_font(doc)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(3.0)
        s.top_margin = s.bottom_margin = Cm(2.5)

    add_cover(doc)
    add_toc(doc)

    i, n = 0, len(lines)
    while i < n:
        raw = lines[i]
        line = raw.rstrip()

        if not line.strip():
            i += 1
            continue

        # 表格
        if line.strip().startswith("|") and i + 1 < n and is_sep(lines[i + 1]):
            header = split_row(line)
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Light Grid Accent 1"
            for j, h in enumerate(header):
                cell = t.rows[0].cells[j]
                cell.text = ""
                rr = cell.paragraphs[0].add_run(h)
                rr.bold = True
                rr.font.size = Pt(10)
                set_ea(rr, "黑体")
                cell.paragraphs[0].paragraph_format.first_line_indent = Pt(0)
            for row in rows:
                cells = t.add_row().cells
                for j in range(min(len(row), len(header))):
                    cells[j].text = ""
                    rr = cells[j].paragraphs[0].add_run(row[j])
                    rr.font.size = Pt(10)
                    set_ea(rr, "宋体")
                    cells[j].paragraphs[0].paragraph_format.first_line_indent = Pt(0)
            doc.add_paragraph()
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            h = doc.add_heading(level=level)
            r = h.add_run(text)
            r.font.color.rgb = RGBColor(0, 0, 0)
            set_ea(r, "黑体")
            sizes = {1: 16, 2: 14, 3: 13, 4: 12}
            r.font.size = Pt(sizes.get(level, 12))
            r.font.bold = True
            h.paragraph_format.first_line_indent = Pt(0)
            i += 1
            continue

        # 引用
        if line.strip().startswith(">"):
            p = doc.add_paragraph()
            add_formatted_runs(p, line.strip().lstrip("> ").strip())
            for r in p.runs:
                r.italic = True
                r.font.size = Pt(10.5)
            p.paragraph_format.left_indent = Pt(20)
            p.paragraph_format.first_line_indent = Pt(0)
            i += 1
            continue

        # 列表
        if re.match(r"^\s*[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            add_formatted_runs(p, re.sub(r"^\s*[-*]\s+", "", line))
            p.paragraph_format.first_line_indent = Pt(0)
            i += 1
            continue
        if re.match(r"^\s*\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            add_formatted_runs(p, re.sub(r"^\s*\d+\.\s+", "", line))
            p.paragraph_format.first_line_indent = Pt(0)
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_formatted_runs(p, line.strip())
        i += 1

    doc.save(OUT)
    print(f"[OK] 已生成 Word 文档 -> {os.path.relpath(OUT, ROOT)}")


if __name__ == "__main__":
    build()
